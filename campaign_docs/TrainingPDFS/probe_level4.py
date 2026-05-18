"""Probe the docx to find Level IV content ranges."""
import docx

doc = docx.Document(r'Original Adventures Reincarnated #3 - Expedition to Barrier Peaks.docx')

# Search for Level IV / Level 4 / Botanical headers
print('=== SEARCHING FOR LEVEL IV MARKERS ===')
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if not text:
        continue
    upper = text.upper()
    if any(kw in upper for kw in ['LEVEL IV', 'LEVEL 4', 'BOTANICAL', 'ROOKERY', 'MENAGERIE', 'GARDEN DECK']):
        print(f'[{i}][{p.style.name}] {text[:120]}')

# Also look for the section headings around levels 3 and 5 to bracket level 4
print('\n=== SEARCHING FOR LEVEL III / V MARKERS ===')
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if not text:
        continue
    upper = text.upper()
    if any(kw in upper for kw in ['LEVEL III', 'LEVEL 3', 'LEVEL V', 'LEVEL 5', 'SERVICE DECK 6', 'WALKWAY AND LOUNGE']):
        print(f'[{i}][{p.style.name}] {text[:120]}')
