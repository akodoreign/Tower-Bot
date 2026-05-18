"""Dump Level IV content from the EBP docx."""
import docx

doc = docx.Document(r'Original Adventures Reincarnated #3 - Expedition to Barrier Peaks.docx')

def dump_range(start, end, label):
    out = []
    for i, p in enumerate(doc.paragraphs[start:end]):
        idx = i + start
        text = p.text.strip()
        if text:
            out.append({'idx': idx, 'style': p.style.name, 'text': text})
    print(f'\n=== {label} ({len(out)} paras) ===')
    for p in out:
        print(f'[{p["idx"]}][{p["style"]}] {p["text"]}')
    return out

# Level IV main content (5e section)
dump_range(2673, 2910, 'LEVEL IV MAIN')

# Chapter 6 Level IV additional encounters
dump_range(3507, 3575, 'LEVEL IV CH6 ADDITIONAL')

# Search for appendix B creatures relevant to Level IV
print('\n=== SEARCHING APPENDIX B FOR LEVEL IV CREATURES ===')
for i, p in enumerate(doc.paragraphs[3884:4800]):
    idx = i + 3884
    text = p.text.strip()
    if not text:
        continue
    upper = text.upper()
    if any(kw in upper for kw in [
        'GASBAT', 'OTYUGH', 'SHAMBLING MOUND', 'VEGEPYGMY', 'TRAPPER',
        'PIERCER', 'GIANT SLUG', 'MIMIC', 'SHOCKER', 'ABOLETH', 'GIANT PYTHON',
        'RHINOCEROS', 'DEINONY', 'OWLBEAR', 'LOPER', 'FLAIL SNAIL',
        'FROGHEMO', 'FROGHEMOTH', 'GIBBERING', 'MOUTHER', 'CATOBLEPAS',
        'BULLYWUG', 'LIZARD KING', 'GIANT LIZARD', 'BASILISK', 'MEDUSA',
        'WILL-O-WISP', 'WISP', 'ROPER', 'KILLER', 'STRANGLE', 'VINE',
        'VAMPIRE THORN', 'OBLIVIAX', 'DARK CREEPER', 'ANHKHEG', 'ANKHEG',
        'UMBER HULK', 'XORN', 'OOZE', 'GELATINOUS', 'OCHRE'
    ]):
        print(f'[{idx}][{p.style.name}] {text[:120]}')

# Appendix C items relevant to Level IV
print('\n=== SEARCHING APPENDIX C FOR LEVEL IV ITEMS ===')
for i, p in enumerate(doc.paragraphs[4798:5090]):
    idx = i + 4798
    text = p.text.strip()
    if not text:
        continue
    upper = text.upper()
    if any(kw in upper for kw in [
        'BINOCULAR', 'HAND GRENADE', 'INCENDIARY', 'PHOSPH', 'SHOCKER',
        'STUN', 'TANGLE', 'TANGLER', 'SLEEP', 'MIRROR', 'HOLOGRAPH',
        'COMMUN', 'RADIO', 'SCANNER', 'DETECT', 'SENSOR', 'DARK',
        'GOGGLES', 'HELMET', 'VISOR', 'JETPACK', 'WING'
    ]):
        print(f'[{idx}][{p.style.name}] {text[:120]}')
