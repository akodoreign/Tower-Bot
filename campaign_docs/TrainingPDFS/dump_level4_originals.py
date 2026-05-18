"""Search original AD&D section for missing creature stats."""
import docx

doc = docx.Document(r'Original Adventures Reincarnated #3 - Expedition to Barrier Peaks.docx')

# Original Level IV is around 529-633
print('=== ORIGINAL LEVEL IV (529-633) ===')
for i, p in enumerate(doc.paragraphs[529:635]):
    idx = i + 529
    text = p.text.strip()
    if text:
        print(f'[{idx}][{p.style.name}] {text[:160]}')

# Also search the full original for leechoid/squealer/snapper
print('\n=== SEARCHING FULL DOC FOR LEECHOID ===')
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if 'LEECHOID' in text.upper():
        print(f'[{i}][{p.style.name}] {text[:160]}')

print('\n=== SEARCHING FULL DOC FOR SQUEALER ===')
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if 'SQUEALER' in text.upper():
        print(f'[{i}][{p.style.name}] {text[:160]}')

print('\n=== SEARCHING FULL DOC FOR SNAPPER-SAW ===')
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if 'SNAPPER' in text.upper():
        print(f'[{i}][{p.style.name}] {text[:160]}')
