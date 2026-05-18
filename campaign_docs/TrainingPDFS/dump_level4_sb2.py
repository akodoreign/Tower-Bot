"""Dump specific stat blocks for Level IV."""
import docx

doc = docx.Document(r'Original Adventures Reincarnated #3 - Expedition to Barrier Peaks.docx')

def dump_range(start, end, label):
    out = []
    for i, p in enumerate(doc.paragraphs[start:end]):
        idx = i + start
        text = p.text.strip()
        if text:
            out.append({'idx': idx, 'style': p.style.name, 'text': text})
    print(f'\n=== {label} ===')
    for p in out:
        print(f'[{p["idx"]}][{p["style"]}] {p["text"][:160]}')

# Key Level IV stat blocks
dump_range(3907, 3935, 'AURUMVORAX')
dump_range(3933, 3955, 'BABOONOID')
dump_range(4076, 4102, 'FLAIL SNAIL')
dump_range(4101, 4125, 'FROGHEMOTH')
dump_range(4121, 4142, 'GLOBE PALM')
dump_range(4186, 4200, 'HORRID PLANT')
dump_range(4224, 4262, 'LIVING BURROW')
dump_range(4249, 4266, 'LIZARDOID')
dump_range(4358, 4374, 'PURPLE BLOSSOM PLANT')
dump_range(4448, 4470, 'SWARM OF ROT GRUBS')
dump_range(4521, 4545, 'TRI-FLOWER FROND')
dump_range(4751, 4775, 'WOLF-IN-SHEEP\'S-CLOTHING')

# Search for LEECHOID and SQUEALER (not found in headings)
print('\n=== SEARCHING FOR LEECHOID ===')
for i, p in enumerate(doc.paragraphs[3884:5090]):
    idx = i + 3884
    text = p.text.strip()
    if 'LEECHOID' in text.upper() or 'LEECH' in text.upper():
        print(f'[{idx}][{p.style.name}] {text[:150]}')

print('\n=== SEARCHING FOR SQUEALER ===')
for i, p in enumerate(doc.paragraphs[3884:5090]):
    idx = i + 3884
    text = p.text.strip()
    if 'SQUEALER' in text.upper():
        print(f'[{idx}][{p.style.name}] {text[:150]}')

print('\n=== SEARCHING FOR SNAPPER-SAW ===')
for i, p in enumerate(doc.paragraphs[3884:5090]):
    idx = i + 3884
    text = p.text.strip()
    if 'SNAPPER' in text.upper():
        print(f'[{idx}][{p.style.name}] {text[:150]}')
