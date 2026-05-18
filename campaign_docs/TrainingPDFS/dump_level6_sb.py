"""Get Level VI specific stat blocks."""
import docx
doc = docx.Document(r'Original Adventures Reincarnated #3 - Expedition to Barrier Peaks.docx')

def dump_range(start, end, label):
    out = []
    for i, p in enumerate(doc.paragraphs[start:end]):
        idx = i + start
        text = p.text.strip()
        if text:
            out.append((idx, p.style.name, text))
    print(f'\n=== {label} ===')
    for idx, style, text in out:
        print(f'[{idx}][{style}] {text[:160]}')

dump_range(3953, 3985, 'BOXING PHYSICAL TRAINING ANDROID')
dump_range(4049, 4080, 'FENCING PHYSICAL TRAINING ANDROID')
dump_range(4198, 4230, 'KARATE PHYSICAL TRAINING ANDROID + BEYOND')
dump_range(4294, 4330, 'PACIFIER ROBOT')
dump_range(4401, 4430, 'SHEDU')
dump_range(4424, 4460, 'STUNTED EYE OF THE DEEP')

# Search for weightlifting android
print('\n=== SEARCHING FOR WEIGHTLIFTING ANDROID ===')
for i, p in enumerate(doc.paragraphs[3884:5090]):
    idx = i + 3884
    text = p.text.strip()
    if 'WEIGHT' in text.upper() or 'LIFTING' in text.upper() or 'FITNESS' in text.upper():
        print(f'[{idx}][{p.style.name}] {text[:160]}')

# Piercers
print('\n=== SEARCHING FOR PIERCER ===')
for i, p in enumerate(doc.paragraphs[3884:5090]):
    idx = i + 3884
    text = p.text.strip()
    if 'PIERCER' in text.upper():
        print(f'[{idx}][{p.style.name}] {text[:160]}')
